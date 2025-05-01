#Q.1 Write a program to create a csv file that we can directly open in MS-Excel.

import csv

# Define the data to write to the CSV
header = ['Name', 'Age', 'Department']
rows = [
    ['Alice', 30, 'HR'],
    ['Bob', 25, 'Engineering'],
    ['Charlie', 28, 'Marketing']
]

# Create and write to a CSV file
with open('employees.csv', mode='w', newline='', encoding='utf-8') as file:
    writer = csv.writer(file)
    
    # Write the header
    writer.writerow(header)
    
    # Write the data rows
    writer.writerows(rows)

print("CSV file 'employees.csv' has been created successfully.")

#Q.2 Read the data stored in MS-Excel file and convert it into a dictionary.
#The record contains rollno, name of student, marks of three subjects. Also calculate total. Display the dictionary data on the monitor.

from openpyxl import load_workbook

# Load the Excel file
wb = load_workbook('students.xlsx')
ws = wb.active

# Dictionary to store student records
students_dict = {}

# Read the data row by row (skip header)
for row in ws.iter_rows(min_row=2, values_only=True):
    rollno, name, sub1, sub2, sub3 = row
    total = sub1 + sub2 + sub3
    
    students_dict[rollno] = {
        'Name': name,
        'Subject1': sub1,
        'Subject2': sub2,
        'Subject3': sub3,
        'Total': total
    }

# Display the dictionary data
for rollno, data in students_dict.items():
    print(f"Roll No: {rollno}")
    for key, value in data.items():
        print(f"  {key}: {value}")
    print()

#Q.3 Accept contact details from the user and create a vcard that we can directly store in our mobile.

    def create_vcard(name, phone, email, org):
    # Basic vCard format (v3.0)
    vcard = f"""BEGIN:VCARD
VERSION:3.0
FN:{name}
TEL;TYPE=CELL:{phone}
EMAIL:{email}
ORG:{org}
END:VCARD
"""
    # Save to file
    filename = f"{name.replace(' ', '_')}.vcf"
    with open(filename, 'w', encoding='utf-8') as file:
        file.write(vcard)

    print(f"vCard saved as '{filename}'. You can now transfer it to your phone.")


# Accept contact details from the user
name = input("Enter full name: ")
phone = input("Enter phone number: ")
email = input("Enter email address: ")
org = input("Enter organization: ")

create_vcard(name, phone, email, org)

#Q.4 Create a specific subdirectory and copy one file from another subdirectory to this newly created subdirectory.

import os
import shutil

# Define paths
source_subdir = 'source_folder'
target_subdir = 'target_folder'
filename = 'example.txt'

# Create full paths
source_path = os.path.join(source_subdir, filename)
target_path = os.path.join(target_subdir, filename)

# Create the target directory if it doesn't exist
os.makedirs(target_subdir, exist_ok=True)

# Copy the file
try:
    shutil.copy2(source_path, target_path)  # copy2 preserves metadata
    print(f"File '{filename}' copied from '{source_subdir}' to '{target_subdir}'.")
except FileNotFoundError:
    print(f"Source file '{source_path}' does not exist.")
except Exception as e:
    print(f"Error copying file: {e}")

#Q.5 Write a program to copy contents of one file to another. While doing so, replace all lowercase characters into uppercase characters.

    def copy_with_uppercase(source_file, destination_file):
    try:
        with open(source_file, 'r', encoding='utf-8') as src:
            content = src.read()

        # Convert all lowercase characters to uppercase
        upper_content = content.upper()

        with open(destination_file, 'w', encoding='utf-8') as dest:
            dest.write(upper_content)

        print(f"Contents copied from '{source_file}' to '{destination_file}' with uppercase transformation.")

    except FileNotFoundError:
        print(f"Error: The file '{source_file}' does not exist.")
    except Exception as e:
        print(f"An error occurred: {e}")


# Example usage
source = 'input.txt'
destination = 'output.txt'
copy_with_uppercase(source, destination)

#Q.6 Write a program that merges lines alternatively from two files and writes the results to new file.
#If one file has less number of lines than the other, the remaining lines from the larger file should be simply copied into the target file.

def merge_alternate_lines(file1, file2, output_file):
    try:
        # Read lines from both files
        with open(file1, 'r', encoding='utf-8') as f1, open(file2, 'r', encoding='utf-8') as f2:
            lines1 = f1.readlines()
            lines2 = f2.readlines()

        # Determine the maximum length
        max_len = max(len(lines1), len(lines2))

        # Merge lines alternately
        merged_lines = []
        for i in range(max_len):
            if i < len(lines1):
                merged_lines.append(lines1[i])
            if i < len(lines2):
                merged_lines.append(lines2[i])

        # Write to output file
        with open(output_file, 'w', encoding='utf-8') as out:
            out.writelines(merged_lines)

        print(f"Lines merged successfully into '{output_file}'.")

    except FileNotFoundError as e:
        print(f"Error: {e.filename} not found.")
    except Exception as e:
        print(f"An error occurred: {e}")


# Example usage
file1 = 'file1.txt'
file2 = 'file2.txt'
output = 'merged_output.txt'
merge_alternate_lines(file1, file2, output)


#Q.7 If an Employee object contains following details:
#empcode, empname, Date of Joining, Salary
#Write a program to serialize and deserialize this data.

import pickle
from datetime import datetime

# Define the Employee class
class Employee:
    def __init__(self, empcode, empname, date_of_joining, salary):
        self.empcode = empcode
        self.empname = empname
        self.date_of_joining = date_of_joining  # Should be a datetime object
        self.salary = salary

    def __str__(self):
        doj = self.date_of_joining.strftime('%Y-%m-%d')
        return f"Employee[{self.empcode}] {self.empname}, Joined: {doj}, Salary: {self.salary}"


# Serialize Employee object to a file
def serialize_employee(emp, filename):
    with open(filename, 'wb') as f:
        pickle.dump(emp, f)
    print(f"Employee object serialized to '{filename}'.")

# Deserialize Employee object from a file
def deserialize_employee(filename):
    with open(filename, 'rb') as f:
        emp = pickle.load(f)
    print(f"Employee object deserialized from '{filename}'.")
    return emp


# Example usage
if __name__ == "__main__":
    # Create an employee
    emp = Employee(
        empcode=101,
        empname="Alice Johnson",
        date_of_joining=datetime(2020, 5, 10),
        salary=75000
    )

    file = 'employee_data.pkl'

    # Serialize
    serialize_employee(emp, file)

    # Deserialize
    loaded_emp = deserialize_employee(file)

    # Display the employee details
    print("\nDeserialized Employee Object:")
    print(loaded_emp)

#Q.8 Given a text file, write a program to create another text file deleting the words ‘a’,
    #‘the’, ‘an’ and replacing each one of them with a blank space.

    import re

def remove_articles(input_file, output_file):
    try:
        with open(input_file, 'r', encoding='utf-8') as infile:
            text = infile.read()

        # Replace whole words 'a', 'an', 'the' (case-insensitive), with a space
        cleaned_text = re.sub(r'\b(a|an|the)\b', ' ', text, flags=re.IGNORECASE)

        # Optionally, remove extra spaces
        cleaned_text = re.sub(r'\s+', ' ', cleaned_text).strip()

        with open(output_file, 'w', encoding='utf-8') as outfile:
            outfile.write(cleaned_text)

        print(f"Articles removed and saved to '{output_file}'.")

    except FileNotFoundError:
        print(f"Error: '{input_file}' not found.")
    except Exception as e:
        print(f"An error occurred: {e}")


# Example usage
input_file = 'original.txt'
output_file = 'cleaned.txt'
remove_articles(input_file, output_file)
